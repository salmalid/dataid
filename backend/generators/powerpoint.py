"""Executive deck. Uses python-pptx when installed; otherwise emits an equivalent
self-contained HTML deck so the artifact always exists."""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Tuple

from . import palette, svg
from ..core.utils import format_metric


def _slides(ctx) -> List[Dict[str, Any]]:
    metrics = ctx.metrics
    if ctx.task_type.value == "classification":
        headline = [("Accuracy", format_metric("accuracy", metrics.get("accuracy"))),
                    ("F1", format_metric("f1", metrics.get("f1"))),
                    ("Recall", format_metric("recall", metrics.get("recall"))),
                    ("ROC AUC", format_metric("roc_auc", metrics.get("roc_auc")))]
    else:
        headline = [("R²", format_metric("r2", metrics.get("r2"))),
                    ("MAE", format_metric("mae", metrics.get("mae"))),
                    ("RMSE", format_metric("rmse", metrics.get("rmse"))),
                    ("MAPE %", format_metric("mape", metrics.get("mape")))]

    top_features = [f["feature"] for f in ctx.explainability.get("global_importance", [])[:5]]
    critical = [i for i in ctx.issues if i.severity == "critical"][:5]
    leaderboard_rows: List[Tuple[str, str]] = []
    if ctx.leaderboard is not None:
        primary = ctx.research.get("primary_metric", "accuracy")
        key = primary if primary in ctx.leaderboard.columns else ctx.leaderboard.columns[-1]
        for _, row in ctx.leaderboard.head(6).iterrows():
            leaderboard_rows.append((str(row["model"]), f"{float(row[key]):.4f}"))

    return [
        {"layout": "title", "title": f"{ctx.target_column}",
         "subtitle": f"dataid automated analysis · run {ctx.run_id} · "
                     f"{datetime.now(timezone.utc).strftime('%d %B %Y')}"},
        {"layout": "bullets", "title": "Executive summary",
         "bullets": [
             f"{ctx.profile.get('rows', 0):,} rows and {ctx.profile.get('n_columns', 0)} columns profiled "
             f"end to end.",
             f"Winning model: {ctx.best_model_name}, selected from "
             f"{0 if ctx.leaderboard is None else len(ctx.leaderboard)} candidates.",
             f"Primary metric: {ctx.research.get('primary_metric')} = "
             f"{headline[0][1]} on data the model never saw.",
             f"Key drivers: {', '.join(top_features[:3]) or 'see explainability'}.",
             "Deployable today as a container with a REST endpoint.",
         ]},
        {"layout": "metrics", "title": "Headline performance", "metrics": headline},
        {"layout": "bullets", "title": "Data quality findings",
         "bullets": [f"{i.severity.upper()} · {i.message}" for i in critical]
         or ["No critical data-quality issues were detected."]},
        {"layout": "table", "title": "Model leaderboard",
         "columns": ["Model", ctx.research.get("primary_metric", "score")],
         "rows": leaderboard_rows},
        {"layout": "bullets", "title": "What drives the prediction",
         "bullets": [f"{i + 1}. {name}" for i, name in enumerate(top_features)]
         or ["Feature attribution unavailable."]},
        {"layout": "bullets", "title": "Deployment",
         "bullets": [
             f"Container image {ctx.deployment.get('image_name', 'dataid/model')}",
             f"p50 latency {ctx.deployment.get('latency', {}).get('p50_ms_per_row', 'n/a')} ms per row",
             f"Targets: {', '.join(ctx.deployment.get('targets', [])[:4])}",
             "REST endpoints: /predict, /predict/batch, /health, /metrics",
         ]},
        {"layout": "bullets", "title": "Recommendations",
         "bullets": _recommendation_bullets(ctx)},
    ]


def _recommendation_bullets(ctx) -> List[str]:
    from .report import _recommendations

    return _recommendations(ctx)[:5]


def generate_powerpoint(ctx) -> Path:
    slides = _slides(ctx)
    try:
        return _pptx(ctx, slides)
    except ImportError:
        return _html_deck(ctx, slides)


def _pptx(ctx, slides: List[Dict[str, Any]]) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    def rgb(hexv: str) -> RGBColor:
        return RGBColor.from_string(hexv.lstrip("#").upper())

    ink = rgb(palette.CHROME["light"]["text_primary"])
    secondary = rgb(palette.CHROME["light"]["text_secondary"])
    accent = rgb(palette.CATEGORICAL_LIGHT[0])

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    for spec in slides:
        slide = presentation.slides.add_slide(blank)

        if spec["layout"] == "title":
            box = slide.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(11.5), Inches(2.2))
            frame = box.text_frame
            frame.text = spec["title"]
            frame.paragraphs[0].runs[0].font.size = Pt(48)
            frame.paragraphs[0].runs[0].font.bold = True
            frame.paragraphs[0].runs[0].font.color.rgb = ink
            para = frame.add_paragraph()
            para.text = spec["subtitle"]
            para.runs[0].font.size = Pt(17)
            para.runs[0].font.color.rgb = secondary
            continue

        header = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.5), Inches(1.0))
        header.text_frame.text = spec["title"]
        run = header.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(30)
        run.font.bold = True
        run.font.color.rgb = ink

        if spec["layout"] == "bullets":
            body = slide.shapes.add_textbox(Inches(0.9), Inches(1.9), Inches(11.3), Inches(4.6))
            frame = body.text_frame
            frame.word_wrap = True
            for i, bullet in enumerate(spec["bullets"]):
                para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                para.text = f"•  {bullet}"
                para.space_after = Pt(14)
                para.runs[0].font.size = Pt(17)
                para.runs[0].font.color.rgb = secondary

        elif spec["layout"] == "metrics":
            for i, (label, value) in enumerate(spec["metrics"]):
                left = Inches(0.9 + i * 3.0)
                card = slide.shapes.add_textbox(left, Inches(2.4), Inches(2.7), Inches(1.8))
                frame = card.text_frame
                frame.text = value
                frame.paragraphs[0].runs[0].font.size = Pt(40)
                frame.paragraphs[0].runs[0].font.bold = True
                frame.paragraphs[0].runs[0].font.color.rgb = accent
                para = frame.add_paragraph()
                para.text = label
                para.runs[0].font.size = Pt(14)
                para.runs[0].font.color.rgb = secondary

        elif spec["layout"] == "table" and spec["rows"]:
            rows, cols = len(spec["rows"]) + 1, len(spec["columns"])
            shape = slide.shapes.add_table(rows, cols, Inches(0.9), Inches(1.9),
                                           Inches(11.3), Inches(0.5 * rows))
            table = shape.table
            for c, name in enumerate(spec["columns"]):
                cell = table.cell(0, c)
                cell.text = str(name)
                cell.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
                cell.text_frame.paragraphs[0].runs[0].font.bold = True
            for r, row in enumerate(spec["rows"], start=1):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = str(value)
                    cell.text_frame.paragraphs[0].runs[0].font.size = Pt(13)

    path = ctx.artifacts_dir / f"deck_{ctx.run_id}.pptx"
    presentation.save(str(path))
    return path


def _html_deck(ctx, slides: List[Dict[str, Any]]) -> Path:
    sections: List[str] = []
    for spec in slides:
        if spec["layout"] == "title":
            sections.append(f"""<section class="slide slide-title">
  <p class="kicker">dataid</p>
  <h1>{svg._esc(spec['title'])}</h1>
  <p class="sub">{svg._esc(spec['subtitle'])}</p>
</section>""")
        elif spec["layout"] == "metrics":
            tiles = "".join(
                f'<div class="metric"><span class="metric-value">{svg._esc(v)}</span>'
                f'<span class="metric-label">{svg._esc(l)}</span></div>'
                for l, v in spec["metrics"])
            sections.append(f'<section class="slide"><h2>{svg._esc(spec["title"])}</h2>'
                            f'<div class="metric-row">{tiles}</div></section>')
        elif spec["layout"] == "table":
            head = "".join(f"<th>{svg._esc(c)}</th>" for c in spec["columns"])
            body = "".join("<tr>" + "".join(f"<td>{svg._esc(c)}</td>" for c in row) + "</tr>"
                           for row in spec["rows"])
            sections.append(f'<section class="slide"><h2>{svg._esc(spec["title"])}</h2>'
                            f'<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></section>')
        else:
            items = "".join(f"<li>{svg._esc(b)}</li>" for b in spec["bullets"])
            sections.append(f'<section class="slide"><h2>{svg._esc(spec["title"])}</h2>'
                            f'<ul>{items}</ul></section>')

    doc = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8" />
<meta name="viewport" content="width=device-width, initial-scale=1" />
<title>{svg._esc(ctx.target_column)} — dataid deck</title>
<style>
{palette.css_variables()}
* {{ box-sizing: border-box; }}
body {{ margin: 0; background: var(--plane); font-family: {svg.FONT};
  color: var(--text-primary); }}
.deck {{ max-width: 1180px; margin: 0 auto; padding: 40px 20px 80px; }}
.slide {{ background: var(--surface-1); border: 1px solid var(--border); border-radius: 18px;
  padding: 56px 60px; margin-bottom: 28px; min-height: 460px; }}
.slide-title {{ display: flex; flex-direction: column; justify-content: center; }}
.kicker {{ font-size: 12px; letter-spacing: 0.16em; text-transform: uppercase;
  color: var(--text-muted); margin: 0 0 18px; }}
h1 {{ font-size: 52px; margin: 0 0 16px; font-weight: 680; line-height: 1.1; }}
h2 {{ font-size: 30px; margin: 0 0 32px; font-weight: 640; }}
.sub, li, td {{ color: var(--text-secondary); font-size: 17px; }}
ul {{ padding-left: 22px; }} li {{ margin-bottom: 16px; line-height: 1.55; }}
.metric-row {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(170px, 1fr)); gap: 24px; }}
.metric {{ display: flex; flex-direction: column; gap: 8px; }}
.metric-value {{ font-size: 42px; font-weight: 680; color: var(--series-1); }}
.metric-label {{ font-size: 14px; color: var(--text-secondary); }}
table {{ width: 100%; border-collapse: collapse; font-variant-numeric: tabular-nums; }}
th, td {{ text-align: left; padding: 12px 14px; border-bottom: 1px solid var(--grid); }}
th {{ color: var(--text-primary); font-size: 14px; }}
@media print {{ .slide {{ page-break-after: always; border: none; }} }}
@media (max-width: 700px) {{ .slide {{ padding: 32px 24px; }} h1 {{ font-size: 34px; }} }}
</style></head>
<body><div class="deck viz-root">{''.join(sections)}</div></body></html>"""

    path = ctx.artifacts_dir / f"deck_{ctx.run_id}.html"
    path.write_text(doc, encoding="utf-8")
    return path
